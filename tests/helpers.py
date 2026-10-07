"""Test doubles and in-memory file builders (no binary fixtures in git)."""

from __future__ import annotations

import io
import re
from typing import Callable, Sequence

import numpy as np

from app.rag.generation import GenerationResult, Passage

# Words in the same group share an embedding dimension, so "car" is a perfect semantic
# match for "automobile". Anything outside every group lands on a residual dimension.
DEFAULT_CONCEPTS = {
    "vehicle": ["car", "automobile", "vehicle", "truck", "engine"],
    "fruit": ["apple", "banana", "fruit", "orchard", "pear"],
    "money": ["refund", "payment", "price", "invoice", "money"],
    "space": ["telescope", "star", "planet", "orbit", "galaxy"],
    "pet": ["dog", "cat", "puppy", "kitten", "pet"],
}


class FakeEmbedder:
    """Deterministic concept-bag embedder. Proves pipeline logic, NOT real retrieval quality."""

    name = "fake-concepts"

    def __init__(self, concepts: dict[str, list[str]] | None = None):
        concepts = concepts or DEFAULT_CONCEPTS
        self._word_to_dim = {w: i for i, words in enumerate(concepts.values()) for w in words}
        self.dimension = len(concepts) + 1  # last dim = "no known concept"
        self.calls: list[list[str]] = []

    def embed(self, texts: Sequence[str]) -> np.ndarray:
        self.calls.append(list(texts))
        out = np.zeros((len(texts), self.dimension), dtype=np.float32)
        for row, text in enumerate(texts):
            known = False
            for word in re.findall(r"[a-z]+", text.lower()):
                dim = self._word_to_dim.get(word)
                if dim is not None:
                    out[row, dim] += 1.0
                    known = True
            if not known:
                out[row, -1] = 1.0
        norms = np.linalg.norm(out, axis=1, keepdims=True)
        return out / norms


def make_pdf(text: str) -> bytes:
    """A minimal single-page PDF whose page content is `text` (Helvetica, one line)."""
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    buf = io.BytesIO()
    buf.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(buf.tell())
        buf.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = buf.tell()
    buf.write(f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n")
    for off in offsets:
        buf.write(f"{off:010d} 00000 n \n".encode())
    buf.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return buf.getvalue()


def make_blank_pdf() -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    from docx import Document

    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class FakeGenerator:
    """Scriptable stand-in for Claude. Records exactly what it was given."""

    available = True

    def __init__(self, reply: GenerationResult | Callable[[str, Sequence[Passage]], GenerationResult] | Exception):
        self._reply = reply
        self.calls: list[tuple[str, list[Passage]]] = []

    def generate(self, question: str, passages: Sequence[Passage]) -> GenerationResult:
        self.calls.append((question, list(passages)))
        if isinstance(self._reply, Exception):
            raise self._reply
        return self._reply(question, passages) if callable(self._reply) else self._reply

    @property
    def seen_texts(self) -> list[str]:
        return [p.text for _, passages in self.calls for p in passages]


def cites(answer: str, *labels: str, grounded: bool = True) -> GenerationResult:
    return GenerationResult(grounded, answer, tuple(labels))


class FakeAnthropicClient:
    """Minimal stand-in for anthropic.Anthropic: records requests, returns a canned reply."""

    def __init__(self, reply_text: str | None = None, *, stop_reason: str = "end_turn", raises: Exception | None = None):
        from types import SimpleNamespace

        self._ns = SimpleNamespace
        self.reply_text, self.stop_reason, self.raises = reply_text, stop_reason, raises
        self.requests: list[dict] = []
        self.messages = self

    def create(self, **kwargs):
        self.requests.append(kwargs)
        if self.raises:
            raise self.raises
        content = [] if self.reply_text is None else [self._ns(type="text", text=self.reply_text)]
        return self._ns(stop_reason=self.stop_reason, content=content)
