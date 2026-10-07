"""Turn uploaded bytes into plain text, dispatching on file extension."""

from __future__ import annotations

import io
import zipfile
from pathlib import PurePath

from app.rag.errors import DocumentLoadError, EmptyDocumentError, UnsupportedDocumentError

TEXT_EXTENSIONS = {".txt", ".md", ".markdown"}
HTML_EXTENSIONS = {".html", ".htm"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | HTML_EXTENSIONS | {".pdf", ".docx"}


def load_text(filename: str, data: bytes) -> str:
    """Extract text from ``data``. Raises a DocumentLoadError subclass on any problem."""
    ext = PurePath(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocumentError(
            f"Unsupported file type {ext or '(none)'!r} for {filename!r}. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if ext in TEXT_EXTENSIONS:
        text = _decode(filename, data)
    elif ext in HTML_EXTENSIONS:
        text = _html_to_text(_decode(filename, data))
    elif ext == ".pdf":
        text = _pdf_to_text(filename, data)
    else:
        text = _docx_to_text(filename, data)

    if not text.strip():
        hint = " (scanned PDFs need OCR, which is not supported)" if ext == ".pdf" else ""
        raise EmptyDocumentError(f"{filename!r} contains no extractable text{hint}")
    return text


def _decode(filename: str, data: bytes) -> str:
    if b"\x00" in data:
        raise DocumentLoadError(f"{filename!r} looks like a binary file, not text")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Windows-1252 covers the common legacy-encoded text files; it never fails except
        # on a few undefined bytes, which we report rather than silently mangle.
        try:
            return data.decode("cp1252")
        except UnicodeDecodeError:
            raise DocumentLoadError(f"{filename!r} is not valid UTF-8 or Windows-1252 text") from None


def _html_to_text(html: str) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript", "template"]):
        tag.decompose()
    return soup.get_text(separator="\n")


def _pdf_to_text(filename: str, data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentLoadError(f"{filename!r} is an encrypted PDF")
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    except DocumentLoadError:
        raise
    except (PyPdfError, ValueError, KeyError, OSError) as exc:
        raise DocumentLoadError(f"{filename!r} could not be read as a PDF: {exc}") from exc


def _docx_to_text(filename: str, data: bytes) -> str:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        doc = Document(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise DocumentLoadError(f"{filename!r} could not be read as a .docx file: {exc}") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n\n".join(parts)
